from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))

BG_COLOR = RGBColor(*hex_to_rgb("#121217"))
TITLE_COLOR = RGBColor(*hex_to_rgb("#6D28D9"))
TEXT_COLOR = RGBColor(*hex_to_rgb("#F3F4F6"))

prs = Presentation()
# Set 16:9 aspect ratio
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Blank slide layout
blank_slide_layout = prs.slide_layouts[6]

slides_data = [
    {
        "title": "QuéVeoHoy\nEstado Actual del Proyecto",
        "content": [
            "Propósito: Gestor y recomendador personal de películas, series y anime.",
            "Stack: Python / CustomTkinter / SQLite",
            "APIs: TMDB y Jikan"
        ],
        "is_title_slide": True
    },
    {
        "title": "2. Pantalla Principal (\"Hoy\")",
        "content": [
            "Tarjeta de recomendación: Destaca el contenido ideal para ver hoy.",
            "Póster centrado: Interfaz atractiva y visual.",
            "Título: Ajuste multilínea automático para nombres largos.",
            "Acciones: Adaptadas dinámicamente por tipo de contenido."
        ]
    },
    {
        "title": "3. Gestión de Series e Importación",
        "content": [
            "Modal Centrado: Selector intuitivo de temporadas y capítulos reales.",
            "Información Precisa: Datos extraídos de las APIs en tiempo real.",
            "Botón Directo: Opción \"Ya la terminé\" para flujo rápido y sin fricción."
        ]
    },
    {
        "title": "4. Pestaña \"En Progreso\"",
        "content": [
            "Seguimiento Episódico: Control exacto de lo que estás viendo.",
            "Tabla Simplificada: Vista clara y sin elementos redundantes.",
            "Reanudación Inmediata: Enlace directo a la vista de \"Hoy\" para continuar viendo."
        ]
    },
    {
        "title": "5. Pestaña \"Biblioteca\"",
        "content": [
            "Catálogo unificado de todo tu contenido.",
            "• Pendientes",
            "• Terminados / Vistos",
            "• Abandonados"
        ]
    },
    {
        "title": "6. Doble Filtro Combinado",
        "content": [
            "Selectores Independientes: Filtra por Tipo (Película, Serie, Anime) y por Estado.",
            "Búsqueda en Tiempo Real: Resultados instantáneos mientras escribes.",
            "Fluidez: La interfaz responde sin bloqueos al filtrar listas grandes."
        ]
    },
    {
        "title": "7. Pestaña \"Historial\"",
        "content": [
            "Auditoría Cronológica: Registro de tus últimas interacciones.",
            "Formato Limpio: Sin horas ni IDs técnicos que ensucien la vista.",
            "Layout Estable: Zona inferior anclada para evitar saltos visuales."
        ]
    },
    {
        "title": "8. Arquitectura Multi-API",
        "content": [
            "Normalización de Anime:",
            "• Unificación de esquemas entre Jikan (Anime) y TMDB (Películas/Series).",
            "• Implementación de bandera interna 'es_anime' para enrutamiento.",
            "• Interfaz unificada a pesar de usar fuentes de datos distintas."
        ]
    },
    {
        "title": "9. Rendimiento y Concurrencia",
        "content": [
            "Buffer de Precarga: Tiempos de respuesta menores a 100ms.",
            "Debounce: Optimización de peticiones en campos de búsqueda.",
            "Request ID Token: Protección contra condiciones de carrera al cambiar de pestaña rápidamente."
        ]
    },
    {
        "title": "10. Estado Actual y Roadmap",
        "content": [
            "Hitos Alcanzados: Sistema base estable, interfaz CustomTkinter, integración TMDB/Jikan.",
            "Próximos Pasos:",
            "• Integración con Letterboxd.",
            "• Calendario semanal de emisiones.",
            "• Animación de ruleta para recomendaciones aleatorias."
        ]
    }
]

for slide_data in slides_data:
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Set background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR
    
    is_title = slide_data.get("is_title_slide", False)
    
    # Add title
    if is_title:
        title_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11.33), Inches(1.5))
    else:
        title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.33), Inches(1))
        
    title_frame = title_box.text_frame
    title_p = title_frame.paragraphs[0]
    title_p.text = slide_data["title"]
    title_p.font.bold = True
    title_p.font.size = Pt(44) if is_title else Pt(40)
    title_p.font.color.rgb = TITLE_COLOR
    if is_title:
        title_p.alignment = PP_ALIGN.CENTER
    
    # Add content box
    if is_title:
        content_box = slide.shapes.add_textbox(Inches(2), Inches(4), Inches(9.33), Inches(3))
    else:
        content_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11.33), Inches(4.5))
        
    content_frame = content_box.text_frame
    content_frame.word_wrap = True
    
    for i, line in enumerate(slide_data["content"]):
        if i == 0:
            p = content_frame.paragraphs[0]
        else:
            p = content_frame.add_paragraph()
        p.text = line
        p.font.size = Pt(28) if is_title else Pt(24)
        p.font.color.rgb = TEXT_COLOR
        if is_title:
            p.alignment = PP_ALIGN.CENTER

prs.save('presentacion_queveohoy.pptx')
print("Archivo presentacion_queveohoy.pptx generado con éxito.")
